import io
import os
import re
import logging
from flask import Flask, render_template, request, send_file, jsonify
from docx import Document
from docxtpl import DocxTemplate
from openpyxl import load_workbook

app = Flask(__name__)
logging.basicConfig(level=logging.INFO)

TEMPLATE_DIR = "server_templates"
os.makedirs(TEMPLATE_DIR, exist_ok=True)

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/api/templates', methods=['GET'])
def get_templates():
    files = os.listdir(TEMPLATE_DIR)
    return jsonify({"docx": [f for f in files if f.endswith('.docx')], "xlsx": [f for f in files if f.endswith('.xlsx')]})

@app.route('/api/upload_template', methods=['POST'])
def upload_template():
    file = request.files['file']
    if file:
        file.save(os.path.join(TEMPLATE_DIR, file.filename))
        return jsonify({"status": "success"})
    return jsonify({"status": "error"}), 400

@app.route('/api/process', methods=['POST'])
def process():
    try:
        template_name = request.form.get('template_name')
        format_type = request.form.get('format')
        action = request.form.get('action')

        if not template_name or not format_type or not action:
            return jsonify({"status": "error", "message": "Missing form data."}), 400

        template_path = os.path.join(TEMPLATE_DIR, os.path.basename(template_name))
        
        if not os.path.exists(template_path):
            return jsonify({"status": "error", "message": f"Template '{template_name}' not found on server."}), 404

        # Read file directly into RAM
        with open(template_path, 'rb') as f:
            stream = io.BytesIO(f.read())

        if format_type == 'docx':
            if action == 'extract':
                doc = Document(stream)
                text = "\n".join([p.text for p in doc.paragraphs])
                for table in doc.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            text += "\n" + cell.text
                variables = list(set(re.findall(r"\{\{(.*?)\}\}", text)))
                return jsonify({"status": "success", "variables": variables})
            
            elif action == 'generate':
                # Get only the variable inputs (ignore format, action, template_name)
                context = {k: v for k, v in request.form.items() if k not in ['format', 'action', 'template_name']}
                
                stream.seek(0)
                doc = DocxTemplate(stream)
                doc.render(context) # If there's a typo in Word, it crashes HERE
                
                output = io.BytesIO()
                doc.save(output)
                output.seek(0)
                return send_file(output, as_attachment=True, download_name=f"Generated_{template_name}", mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')

        elif format_type == 'xlsx':
            if action == 'extract':
                wb = load_workbook(stream)
                text = ""
                for sheet in wb.worksheets:
                    for row in sheet.iter_rows():
                        for cell in row:
                            if cell.value: text += str(cell.value) + "\n"
                variables = list(set(re.findall(r"\{\{(.*?)\}\}", text)))
                return jsonify({"status": "success", "variables": variables})
            
            elif action == 'generate':
                context = {k: v for k, v in request.form.items() if k not in ['format', 'action', 'template_name']}
                stream.seek(0)
                wb = load_workbook(stream)
                for sheet in wb.worksheets:
                    for row in sheet.iter_rows():
                        for cell in row:
                            if cell.value and isinstance(cell.value, str):
                                for key, value in context.items():
                                    cell.value = cell.value.replace(f"{{{{{key}}}}}", str(value))
                
                output = io.BytesIO()
                wb.save(output)
                output.seek(0)
                return send_file(output, as_attachment=True, download_name=f"Generated_{template_name}", mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

        return jsonify({"status": "error", "message": "Invalid action or format."}), 400

    except Exception as e:
        # This catches ANY crash (like jinja2 template errors) and sends it to the screen
        app.logger.error(f"CRASH: {str(e)}")
        return jsonify({"status": "error", "message": f"Server Error: {str(e)}"}), 500

if __name__ == '__main__':
    app.run(debug=True, host='0.0.0.0', port=5000)